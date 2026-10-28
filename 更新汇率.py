# -*- coding: utf-8 -*-
"""
更新汇率.py — 每日自动获取 USD→CNY 写入成本表「参数!B2」
用法：与成本表放同一文件夹，双击或：python 更新汇率.py
"""
from __future__ import annotations

import json
import os
import sys
import urllib.request
from datetime import date
from pathlib import Path

try:
    import openpyxl
except ImportError:
    os.system(f'"{sys.executable}" -m pip install openpyxl -q')
    import openpyxl

HERE = Path(__file__).resolve().parent
COST_NAMES = [
    "成本分析模板_副本2_已修改.xlsx",
    "成本分析模板_副本2_已修改.xlsm",
]


def fetch_usd_cny():
    urls = [
        "https://open.er-api.com/v6/latest/USD",
        "https://api.exchangerate-api.com/v4/latest/USD",
    ]
    last = None
    for url in urls:
        try:
            req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(req, timeout=15) as r:
                d = json.load(r)
            cny = d.get("rates", {}).get("CNY")
            if cny:
                asof = d.get("time_last_update_utc") or d.get("date") or date.today().isoformat()
                return float(cny), str(asof), url
        except Exception as e:
            last = e
    raise RuntimeError(f"获取汇率失败: {last}")


def main():
    cost = next((HERE / n for n in COST_NAMES if (HERE / n).exists()), None)
    if not cost:
        print("未找到成本表，请与本脚本放同一文件夹")
        input("按回车退出")
        return
    print("正在获取 USD→CNY …")
    rate, asof, src = fetch_usd_cny()
    print(f"今日汇率: {rate}  ({asof})")
    wb = openpyxl.load_workbook(cost)
    if "参数" not in wb.sheetnames:
        ws = wb.create_sheet("参数", 0)
        ws["A2"] = "美元兑人民币汇率 (USD→CNY)"
    else:
        ws = wb["参数"]
    ws["B2"] = rate
    ws["B2"].number_format = "0.0000"
    ws["B3"] = asof
    ws["A3"] = ws["A3"].value or "更新日期/来源时间"
    ws["B4"] = src
    ws["A4"] = ws["A4"].value or "数据来源"
    # 确保主表 M 列引用参数
    if "采购成本对比" in wb.sheetnames:
        buy = wb["采购成本对比"]
        for r in range(4, 21):
            buy.cell(r, 13, "=参数!$B$2")
            buy.cell(r, 13).number_format = "0.00"
    wb.save(cost)
    print(f"已写入 {cost.name} → 参数!B2={rate}")
    print("主表「汇率」列已引用该值。重新打开表格即可。")
    input("按回车退出")


if __name__ == "__main__":
    try:
        main()
    except Exception as e:
        print("失败:", e)
        import traceback
        traceback.print_exc()
        input("按回车退出")
